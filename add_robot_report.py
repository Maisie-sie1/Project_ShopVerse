"""เพิ่ม Robot Framework results ลง test-automation.xlsx (สร้าง sheet Robot Tests)"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FN = "test-automation.xlsx"
wb = openpyxl.load_workbook(FN)

HEADER_FONT = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", start_color="1E3A8A")
BOLD = Font(name="Segoe UI", size=10, bold=True)
REGULAR = Font(name="Segoe UI", size=10)
PASS_FILL = PatternFill("solid", start_color="C6EFCE")
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

if "Robot Tests" in wb.sheetnames:
    del wb["Robot Tests"]
ws = wb.create_sheet("Robot Tests")

headers = ["ID", "Feature", "Scenario", "Framework", "Result", "Bug ID", "Notes"]
ws.append(headers)
for c in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=c)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    cell.border = BORDER

rows = [
    ("RF-STORE-001", "Storefront", "หน้าแรก + สินค้าแนะนำ", "Robot+Selenium", "PASS", "", ""),
    ("RF-STORE-002", "Storefront", "CTA → /shop", "Robot+Selenium", "PASS", "", ""),
    ("RF-STORE-003", "Storefront", "เปิดสินค้า + รายละเอียด", "Robot+Selenium", "PASS", "", ""),
    ("RF-STORE-004", "Storefront", "ค้นหา กาแฟ", "Robot+Selenium", "PASS", "", ""),
    ("RF-AUTH-001", "Auth", "สมัครสมาชิก + เข้าสู่ระบบ", "Robot+Selenium", "PASS", "", ""),
    ("RF-AUTH-002", "Auth", "สมัครรหัสไม่ตรง → error", "Robot+Selenium", "PASS", "", ""),
    ("RF-AUTH-003", "Auth", "login seed customer", "Robot+Selenium", "PASS", "", ""),
    ("RF-AUTH-004", "Auth", "login รหัสผิด → error", "Robot+Selenium", "PASS", "", ""),
    ("RF-AUTH-005", "Auth", "logout ผ่าน UI", "Robot+Selenium", "PASS", "BUG-ROBOT-001", "ใช้ JS click workaround (native click ไม่เปิด dropdown)"),
    ("RF-AUTH-006", "Auth/RBAC", "customer → /admin redirect", "Robot+Selenium", "PASS", "", ""),
    ("RF-CART-001", "Cart", "เพิ่มสินค้า → badge", "Robot+Selenium", "PASS", "", ""),
    ("RF-CART-002", "Cart", "/cart เปลี่ยนจำนวน", "Robot+Selenium", "PASS", "", ""),
    ("RF-CART-003", "Cart", "ลบสินค้า → ว่าง", "Robot+Selenium", "PASS", "", ""),
    ("RF-CHK-001", "Checkout", "สั่งซื้อ COD สำเร็จ → ORD-", "Robot+Selenium", "PASS", "", ""),
    ("RF-CHK-002", "Checkout", "ไม่กรอกที่อยู่ → validation", "Robot+Selenium", "PASS", "", ""),
    ("RF-API-001", "API/Auth", "register 201 ไม่ leak hash", "Robot+Requests", "PASS", "", ""),
    ("RF-API-002", "API/Auth", "register ซ้ำ 409", "Robot+Requests", "PASS", "", ""),
    ("RF-API-003", "API/Auth", "login 200", "Robot+Requests", "PASS", "", ""),
    ("RF-API-004", "API/Auth", "login รหัสผิด 401", "Robot+Requests", "PASS", "", ""),
    ("RF-API-005", "API/Products", "GET products + pagination", "Robot+Requests", "PASS", "", ""),
    ("RF-API-006", "API/Products", "กรองหมวด drinks", "Robot+Requests", "PASS", "", ""),
    ("RF-API-007", "API/Security", "GET cart ไม่ login → 401", "Robot+Requests", "PASS", "", ""),
    ("RF-API-008", "API/RBAC", "customer → admin/stats 403", "Robot+Requests", "PASS", "", ""),
]
for r, row in enumerate(rows, start=2):
    ws.append(row)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=r, column=c)
        cell.font = REGULAR
        cell.alignment = WRAP
        cell.border = BORDER
    if row[4] == "PASS":
        ws.cell(row=r, column=5).fill = PASS_FILL

for i, w in enumerate([16, 16, 30, 18, 10, 14, 40], start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

# อัปเดต Summary ให้มี Robot totals
ws_s = wb["Summary"]
summary_add = [
    ("Robot Framework Tests", 23),
    ("Robot PASS", 23),
    ("Robot FAIL", 0),
    ("Robot Bugs Found", 1),
]
start = ws_s.max_row + 2
for i, (k, v) in enumerate(summary_add):
    ws_s.cell(row=start + i, column=1, value=k).font = BOLD
    ws_s.cell(row=start + i, column=2, value=v).font = REGULAR

wb.save(FN)
print("✅ Added 'Robot Tests' sheet + summary to", FN)